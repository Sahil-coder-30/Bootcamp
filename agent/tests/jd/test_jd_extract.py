"""Role 1 JD Analytics: sections, evidence gate, levels, company, pipeline, API. Offline (fake LLM)."""

import base64
import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from talent_match.documents import parse_document, parse_text
from talent_match.jd import pipeline as jd_pipeline
from talent_match.jd.company import CompanyRef, resolve_company
from talent_match.jd.levels import detect_seniority, phrase_level, years_level
from talent_match.jd.pipeline import JdExtractRequest, extract_jd, line_starts, prepare
from talent_match.jd.postprocess import EvidenceIndex, split_names
from talent_match.jd.schema import JdLlmOutput, LlmSkill
from talent_match.jd.sections import match_heading, structure
from talent_match.llm import AllProvidersFailed, LlmResult, PromptSpec, load_prompt, run_structured

FIX = Path(__file__).resolve().parents[1] / "fixtures" / "jds"
GOOGLE = (FIX / "google_swe.txt").read_text()
ORACLE = (FIX / "oracle_fss_support.txt").read_text()
COMPANIES = [
    CompanyRef(company_id=4, name="Google LLC"),
    CompanyRef(company_id=78, name="Microsoft Corporation"),
    CompanyRef(company_id=24, name="Oracle Financial Services Software"),
    CompanyRef(company_id=3, name="Oracle India Pvt Ltd"),
    CompanyRef(company_id=9, name="Zomato Ltd"),
]


def fake_llm(output: JdLlmOutput):
    def run(schema, prompt, variables, **_):
        assert "L001" in variables["numbered_jd"]
        return LlmResult(value=output, provider="fake", model="fake-1", latency_ms=5)

    return run


GOOGLE_LLM = JdLlmOutput(
    company="Google",
    role="Software Engineer",
    seniority="fresher",
    skills=[
        LlmSkill(skill_name="Java", category_code="COD", evidence="such as Java, C++, Python or Go", line=8),
        LlmSkill(
            skill_name="Data structures and algorithms",
            category_code="DSA",
            evidence="Strong knowledge of data structures and algorithms",
            line=9,
            strength="strong",
        ),
        LlmSkill(
            skill_name="Kubernetes",
            category_code="CLOUD",
            evidence="Familiarity with Kubernetes",
            line=14,
            requirement="must",
        ),
        LlmSkill(
            skill_name="Large-scale systems",
            category_code="SYSD",
            explicit=False,
            evidence="building systems that serve millions of users",
            line=13,
        ),
        LlmSkill(skill_name="Rust", category_code="COD", evidence="Expert Rust programmer for embedded work"),
        LlmSkill(
            skill_name="Python", category_code="COD", evidence="paraphrased python mention that is not there"
        ),
    ],
)


# --- documents + sections -------------------------------------------------------


def test_headings_map_to_kinds():
    assert match_heading("What We're Looking For") == ("requirements", "must", False)
    assert match_heading("Nice to Have")[1] == "nice"
    assert match_heading("Benefits")[2] is True
    assert match_heading("Python, Java and Go are used daily here") is None


def test_structure_sections_and_offsets():
    doc = structure(parse_text(GOOGLE))
    kinds = [(s.kind, s.requirement_hint, s.drop) for s in doc.sections]
    assert ("requirements", "must", False) in kinds
    assert ("requirements", "nice", False) in kinds
    assert ("other", None, True) in kinds  # Benefits
    req = next(s for s in doc.sections if s.requirement_hint == "must" and s.kind == "requirements")
    assert doc.text[req.start : req.end].startswith("Minimum Qualifications")
    assert doc.pages == [(0, len(doc.text))]


def test_prepare_drops_boilerplate_and_tags_sections():
    doc = structure(parse_text(GOOGLE))
    p = prepare(doc)
    assert '[SECTION requirements | nice | "Preferred Qualifications"]' in p
    assert "Free meals" not in p
    assert p.splitlines()[0].startswith("L001 Google LLC")


def test_docx_and_pdf_parse():
    import docx
    import pymupdf

    d = docx.Document()
    d.add_heading("Requirements", level=1)
    d.add_paragraph("Strong knowledge of Python and SQL.", style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    doc = structure(parse_document(buf.getvalue(), "jd.docx"))
    assert any(s.kind == "requirements" for s in doc.sections)

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Requirements", fontsize=16)
    page.insert_text(
        (72, 100), "- Hands-on experience with Docker and Kubernetes in production.", fontsize=10
    )
    doc = structure(parse_document(pdf.tobytes(), "jd.pdf"))
    assert "Kubernetes" in doc.text
    assert any(s.kind == "requirements" for s in doc.sections)


# --- evidence gate + names ----------------------------------------------------------


def test_evidence_gate_exact_fuzzy_rebuilt_dropped():
    text = "Requirements\n- Strong knowledge of data structures and algorithms.\n- Experience with SQL."
    idx = EvidenceIndex(text, line_starts(text))
    exact = idx.locate("strong knowledge of DATA structures", 2, "DSA")
    assert exact.method == "verified_exact"
    assert text[exact.start : exact.end] == "Strong knowledge of data structures"
    fuzzy = idx.locate("Strong knowledge of data structure and algorithm", None, "DSA")
    assert fuzzy.method in ("verified_fuzzy", "weak")
    rebuilt = idx.locate("totally different words here", None, "SQL")
    assert rebuilt.method == "rebuilt" and "SQL" in text[rebuilt.start : rebuilt.end]
    assert idx.locate("nothing like this exists", None, "Haskell") is None


def test_split_names_only_into_distinct_known_skills():
    assert split_names("Java/C++") == ["Java", "C++"]
    assert split_names("Data structures and algorithms") == ["Data structures and algorithms"]
    assert split_names("Kafka or FooBarBaz") == ["Kafka or FooBarBaz"]


# --- levels + company -----------------------------------------------------------------


def test_level_rules():
    assert phrase_level("Strong knowledge of Java")[0] == 5
    assert phrase_level("Excellent hands-on Java")[0] == 6
    assert phrase_level("Exposure to Docker")[0] == 2
    assert phrase_level("Architect our platform")[0] == 9
    assert years_level("3+ years of Go")[0] == 7
    assert years_level("0-1 years")[0] == 4
    assert detect_seniority("Senior Software Engineer", None, None)[0] == "senior"
    assert detect_seniority("Software Engineer", None, "junior") == ("junior", "llm")
    assert detect_seniority("Engineer", None, None) == ("mid", "default")


def test_company_resolution():
    assert resolve_company("Google LLC", "", COMPANIES)[:2] == (4, "Google LLC")
    assert resolve_company("Alphabet", "", COMPANIES)[0] == 4
    assert resolve_company("Microsoft", "", [])[0] == 78
    assert resolve_company("Oracle", "Oracle ... financial services industry", COMPANIES)[0] == 24
    assert resolve_company("OFSS", "", COMPANIES)[0] == 24
    assert resolve_company("Zomato", "", COMPANIES)[:3] == (9, "Zomato Ltd", "fuzzy")
    assert resolve_company("Unknown Startup", "", COMPANIES)[0] is None
    assert resolve_company(None, "", COMPANIES, company_id=78)[0] == 78


# --- pipeline ----------------------------------------------------------------------------


def test_pipeline_with_llm(monkeypatch):
    monkeypatch.setattr(jd_pipeline, "run_structured", fake_llm(GOOGLE_LLM))
    out = extract_jd(JdExtractRequest(text=GOOGLE, source_file="google_swe.txt", companies=COMPANIES))
    by = {s.skill_name: s for s in out.skills}
    assert out.company_id == 4 and out.company == "Google LLC"
    assert out.meta["model"] == "fake:fake-1" and out.meta["degraded"] is False
    assert by["Java"].confidence == "high" and by["Java"].requirement == "must"
    assert by["Data structures and algorithms"].canonical_id == "dsa"
    # LLM said must, but the line sits under "Preferred Qualifications"
    assert by["Kubernetes"].requirement == "nice"
    assert by["Large-scale systems"].confidence == "low"
    # Rust has no evidence and is not in the JD: dropped by the gate
    assert "Rust" not in by
    assert any(d["skill_name"] == "Rust" for d in out.meta["dropped"])
    # Python evidence is invented, but "Python" is in the text: evidence rebuilt from the sentence
    assert "Python" in by["Python"].evidence and by["Python"].confidence == "medium"
    assert set(out.category_levels) <= {"COD", "DSA", "CLOUD", "SYSD", "SQL", "AI", "SWE", "COMM"}
    assert 3 <= out.category_levels["DSA"] <= 6  # fresher band
    assert out.meta["level_basis"]["DSA"]


def test_evidence_spans_point_into_text(monkeypatch):
    monkeypatch.setattr(jd_pipeline, "run_structured", fake_llm(GOOGLE_LLM))
    doc = structure(parse_text(GOOGLE))
    out = extract_jd(JdExtractRequest(document=doc, source_file="g.txt"))
    for s in out.skills:
        assert doc.text[s.evidence_span["start"] : s.evidence_span["end"]].strip() == s.evidence


def test_pipeline_degraded_when_all_providers_fail(monkeypatch):
    def boom(*_a, **_k):
        raise AllProvidersFailed([{"provider": "x", "error": "429", "kind": "rate_limit"}])

    monkeypatch.setattr(jd_pipeline, "run_structured", boom)
    out = extract_jd(JdExtractRequest(text=ORACLE, source_file="oracle.txt", companies=COMPANIES))
    assert out.meta["degraded"] is True
    assert "degraded:all_llm_providers_failed" in out.meta["warnings"]
    assert out.company_id == 24
    ids = {s.canonical_id for s in out.skills}
    assert {"sql", "java", "linux"} <= ids
    assert all(s.confidence == "low" for s in out.skills)


def test_keyword_only_mode():
    out = extract_jd(JdExtractRequest(text=GOOGLE, use_llm=False))
    assert out.meta["degraded"] and out.skills
    assert out.meta["prompt_version"] == "jd-v1"


# --- llm.py -------------------------------------------------------------------------------


class _Msg:
    def __init__(self, content):
        self.content = content


class _FakeModel:
    def __init__(self, answers):
        self.answers = list(answers)
        self.model_name = "fake-model"

    def invoke(self, _messages):
        a = self.answers.pop(0)
        if isinstance(a, Exception):
            raise a
        return _Msg(a)


def test_run_structured_repair_and_fallback(monkeypatch):
    from talent_match import llm as llm_mod

    good = '{"skills": [{"skill_name": "Go", "category_code": "cod", "evidence": "Go", "line": "L003"}]}'
    models = {"a": _FakeModel([TimeoutError("timed out")]), "b": _FakeModel(["not json", good])}
    monkeypatch.setattr(llm_mod, "_model", lambda name, _t: models[name])
    spec = PromptSpec(name="t", version="t-1", system="sys", user="hi {{x}}")
    res = run_structured(JdLlmOutput, spec, {"x": "1"}, provider_order=["a", "b"])
    assert (
        res.provider == "b" and res.value.skills[0].category_code == "COD" and res.value.skills[0].line == 3
    )
    assert [a["kind"] for a in res.attempts] == ["timeout", "invalid"]

    models2 = {"a": _FakeModel(["{}", "{}"])}
    monkeypatch.setattr(llm_mod, "_model", lambda name, _t: models2[name])
    with pytest.raises(AllProvidersFailed):
        run_structured(JdLlmOutput, spec, {"x": "1"}, provider_order=["a"])


def test_prompt_loads_with_version():
    spec = load_prompt("jd_extract")
    assert spec.version == "jd-v1"
    sys_, user = spec.render(
        {"source_file": "a.pdf", "company_hint": "-", "role_hint": "-", "numbered_jd": "L001 x"}
    )
    assert "EXAMPLE" in sys_ and "L001 x" in user and "{{" not in user


# --- API -------------------------------------------------------------------------------------


@pytest.fixture
def client():
    from openradix_agents.api.main import create_app

    with TestClient(create_app()) as c:
        yield c


def test_api_documents_parse_and_extract(client, monkeypatch):
    b64 = base64.b64encode(GOOGLE.encode()).decode()
    r = client.post("/documents/parse", json={"file_name": "google.txt", "content_b64": b64})
    assert r.status_code == 200
    doc = r.json()
    assert len(doc["sha256"]) == 64 and doc["sections"]

    monkeypatch.setattr(jd_pipeline, "run_structured", fake_llm(GOOGLE_LLM))
    doc.pop("sha256"), doc.pop("file_name"), doc.pop("size_bytes")
    r = client.post(
        "/extract/jd",
        json={
            "document": doc,
            "source_file": "google.txt",
            "document_id": "d1",
            "companies": [c.model_dump() for c in COMPANIES],
        },
    )
    assert r.status_code == 200
    body = r.json()
    assert body["source_type"] == "jd" and body["company_id"] == 4 and body["document_id"] == "d1"


def test_api_errors(client):
    assert (
        client.post("/documents/parse", json={"file_name": "a.txt", "content_b64": "%%%"}).status_code == 400
    )
    empty = base64.b64encode(b"").decode()
    assert (
        client.post("/documents/parse", json={"file_name": "a.txt", "content_b64": empty}).status_code == 422
    )
    assert client.post("/extract/jd", json={"source_file": "x"}).status_code == 400


def test_api_version(client):
    body = client.get("/extract/jd/version").json()
    assert body["prompt_version"] == "jd-v1" and body["parser_version"] == "doc-1.0"
