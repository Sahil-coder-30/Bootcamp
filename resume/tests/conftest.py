"""Shared fixtures: sample resumes rendered as PDF (single and two-column) and DOCX,
and an app wired to a temp SQLite DB + local file storage. Nothing touches the network."""

from __future__ import annotations

import io
from collections.abc import Iterator
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

FIXTURES = Path(__file__).parent / "fixtures"
CANDIDATES = ("ananya_rao", "karthik_subramaniam", "priya_menon", "rohan_verma")


def resume_text(name: str) -> str:
    return (FIXTURES / f"{name}.txt").read_text(encoding="utf-8")


def _is_heading(line: str) -> bool:
    from resume_service.sections import classify_heading

    s = line.strip()
    return bool(s) and len(s.split()) <= 4 and ":" not in s and classify_heading(s) is not None


def make_pdf(text: str) -> bytes:
    """Single-column PDF: headings bold and larger, like a real resume template."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    _w, h = A4
    y = h - 50
    for i, line in enumerate(text.splitlines()):
        if not line.strip():
            y -= 6
            continue
        if i == 0:
            c.setFont("Helvetica-Bold", 18)
        elif _is_heading(line):
            y -= 4
            c.setFont("Helvetica-Bold", 13)
        else:
            c.setFont("Helvetica", 10)
        for chunk in _wrap(line, 105):
            if y < 50:
                c.showPage()
                y = h - 50
                c.setFont("Helvetica", 10)
            c.drawString(40, y, ("• " + chunk[2:]) if chunk.startswith("- ") else chunk)
            y -= 14
    c.save()
    return buf.getvalue()


def make_two_column_pdf(sidebar: list[str], main: list[str], name: str) -> bytes:
    """Canva-style layout: a left sidebar (contact + skills) and a main column."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    _w, h = A4
    c.setFont("Helvetica-Bold", 20)
    c.drawString(40, h - 50, name)
    for x, lines, width in ((40, sidebar, 30), (230, main, 70)):
        y = h - 90
        for line in lines:
            bold = _is_heading(line)
            c.setFont("Helvetica-Bold" if bold else "Helvetica", 12 if bold else 9.5)
            for chunk in _wrap(line, width):
                c.drawString(x, y, chunk)
                y -= 13
            if bold:
                y -= 2
    c.save()
    return buf.getvalue()


def make_docx(text: str, header_line: str | None = None, table_layout: bool = False) -> bytes:
    import docx

    d = docx.Document()
    lines = text.splitlines()
    if header_line:
        d.sections[0].header.paragraphs[0].text = header_line
    target = d
    if table_layout:
        # Two-cell layout table: name block on the left, the rest on the right.
        tbl = d.add_table(rows=1, cols=2)
        tbl.cell(0, 0).paragraphs[0].text = lines[0]
        target = tbl.cell(0, 1)
        lines = lines[1:]
    for line in lines:
        if not line.strip():
            continue
        if _is_heading(line):
            p = target.add_paragraph()
            run = p.add_run(line.strip())
            run.bold = True
        elif line.startswith("- "):
            target.add_paragraph(line[2:], style="List Bullet")
        else:
            target.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _wrap(line: str, width: int) -> list[str]:
    words, out, cur = line.split(" "), [], ""
    for w in words:
        if len(cur) + len(w) + 1 > width and cur:
            out.append(cur)
            cur = "  " + w
        else:
            cur = f"{cur} {w}".strip() if not cur.startswith("  ") else f"{cur} {w}"
    if cur:
        out.append(cur)
    return out or [""]


@pytest.fixture(scope="session")
def texts() -> dict[str, str]:
    return {n: resume_text(n) for n in CANDIDATES}


@pytest.fixture()
def settings(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from resume_service.settings import Settings

    for var in ("S3_BUCKET", "RESUME_S3_BUCKET", "RESUME_API_KEY", "DATABASE_URL", "RESUME_DATABASE_URL"):
        monkeypatch.delenv(var, raising=False)
    return Settings(database_url=f"sqlite:///{tmp_path / 'resumes.db'}", storage_dir=tmp_path / "files",
                    max_upload_mb=1.0)  # type: ignore[call-arg]


@pytest.fixture()
def client(settings) -> Iterator[TestClient]:
    from resume_service.main import create_app

    with TestClient(create_app(settings)) as c:
        yield c
