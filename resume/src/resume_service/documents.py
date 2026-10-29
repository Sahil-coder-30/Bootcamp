"""
documents.py
============
PDF / DOCX / TXT bytes -> ordered ``Line`` objects with the layout hints that
section detection needs (font size, bold, caps). No network, no LLM.

PDF: PyMuPDF spans -> lines -> blocks, two-column detection with an x-occupancy
histogram (plan 3.2), running header/footer removal, private-use icon glyphs and
ligatures cleaned. DOCX: body paragraphs and tables in document order, text boxes,
and the first section header (contact details often live there).
"""

from __future__ import annotations

import io
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass, field

MAX_PAGES = 15


class DocumentError(ValueError):
    """The bytes could not be read as a resume. ``code`` is a stable machine-readable id."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass
class Line:
    text: str
    page: int = 0
    x0: float = 0.0
    y0: float = 0.0
    size: float = 11.0
    bold: bool = False
    style: str | None = None  # DOCX paragraph style name
    region: str = "body"  # body | header


@dataclass
class ParsedDocument:
    lines: list[Line]
    page_count: int
    layout: str = "single"  # single | two_column | docx | text
    warnings: list[str] = field(default_factory=list)
    ocr_used: bool = False

    @property
    def text(self) -> str:
        return "\n".join(line.text for line in self.lines)


_BULLETS = re.compile(r"^[\s]*[•▪●➢◦‣∙·■□►▶✓✔\-–—*]+\s*")
_PUA = re.compile("[-]")


def clean_text(s: str) -> str:
    s = unicodedata.normalize("NFKC", s)
    s = _PUA.sub(" ", s)
    s = s.replace(" ", " ").replace("\t", " ")
    if _BULLETS.match(s) and not re.match(r"^\s*-?\d", s):
        s = "- " + _BULLETS.sub("", s, count=1)
    return re.sub(r"[ ]{2,}", " ", s).strip()


def sniff_kind(data: bytes, file_name: str, mime: str | None = None) -> str:
    name = file_name.lower()
    if data[:5] == b"%PDF-" or name.endswith(".pdf") or mime == "application/pdf":
        return "pdf"
    if data[:2] == b"PK" and (name.endswith(".docx") or not name.endswith(".zip")):
        return "docx"
    if name.endswith((".txt", ".md")) or (mime or "").startswith("text/"):
        return "text"
    try:
        data[:4096].decode("utf-8")
        return "text"
    except UnicodeDecodeError as exc:
        raise DocumentError("unsupported_format", "Upload a PDF, DOCX or TXT resume.") from exc


def parse_document(data: bytes, file_name: str, mime: str | None = None) -> ParsedDocument:
    if not data:
        raise DocumentError("empty_file", "The uploaded file is empty.")
    kind = sniff_kind(data, file_name, mime)
    if kind == "pdf":
        doc = _parse_pdf(data)
    elif kind == "docx":
        doc = _parse_docx(data)
    else:
        doc = _parse_text(data.decode("utf-8", errors="replace"))
    doc.lines = _join_hyphenated([ln for ln in doc.lines if ln.text])
    if sum(len(ln.text) for ln in doc.lines) < 40:
        if kind == "pdf":
            raise DocumentError(
                "scanned_pdf_ocr_unavailable",
                "This PDF has no text layer (it looks scanned). Upload the original PDF or a DOCX.")
        raise DocumentError("unreadable_document", "Could not find any text in this document.")
    return doc


def parse_text(text: str) -> ParsedDocument:
    doc = _parse_text(text)
    doc.lines = _join_hyphenated([ln for ln in doc.lines if ln.text])
    return doc


# --- plain text ---------------------------------------------------------------

def _parse_text(text: str) -> ParsedDocument:
    lines = [Line(text=clean_text(t)) for t in text.splitlines()]
    return ParsedDocument(lines=lines, page_count=1, layout="text")


def _join_hyphenated(lines: list[Line]) -> list[Line]:
    out: list[Line] = []
    for ln in lines:
        if out and out[-1].text.endswith("-") and not out[-1].text.endswith(" -") and ln.text[:1].islower():
            out[-1].text = out[-1].text[:-1] + ln.text
            continue
        out.append(ln)
    return out


# --- PDF ----------------------------------------------------------------------

@dataclass
class _Block:
    lines: list[Line]
    x0: float
    y0: float
    x1: float
    y1: float

    @property
    def width(self) -> float:
        return self.x1 - self.x0

    @property
    def height(self) -> float:
        return max(1.0, self.y1 - self.y0)


def _parse_pdf(data: bytes) -> ParsedDocument:
    import pymupdf

    try:
        pdf = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:  # corrupt bytes raise a variety of errors
        raise DocumentError("unreadable_document", f"Could not open the PDF: {exc}") from exc
    if pdf.needs_pass:
        raise DocumentError("unreadable_document", "The PDF is password protected.")

    warnings: list[str] = []
    page_blocks: list[tuple[int, float, float, list[_Block]]] = []
    for pno, page in enumerate(pdf):
        if pno >= MAX_PAGES:
            warnings.append(f"truncated_after_{MAX_PAGES}_pages")
            break
        blocks = _page_blocks(page.get_text("dict"), pno)
        page_blocks.append((pno, page.rect.width, page.rect.height, blocks))

    _drop_running_headers(page_blocks)
    layout = "single"
    lines: list[Line] = []
    for _pno, width, height, blocks in page_blocks:
        ordered, two_col = _reading_order(blocks, width, height)
        if two_col:
            layout = "two_column"
        for b in ordered:
            lines.extend(b.lines)
    return ParsedDocument(lines=lines, page_count=len(pdf), layout=layout, warnings=warnings)


def _page_blocks(d: dict, pno: int) -> list[_Block]:
    blocks: list[_Block] = []
    for blk in d.get("blocks", []):
        if blk.get("type") != 0:
            continue
        lines: list[Line] = []
        for ln in blk.get("lines", []):
            spans = [s for s in ln.get("spans", []) if s.get("text", "").strip()]
            if not spans:
                continue
            text = clean_text("".join(s["text"] for s in ln["spans"]))
            if not text:
                continue
            size = max(s["size"] for s in spans)
            bold = all(("bold" in s.get("font", "").lower()) or (s.get("flags", 0) & 16) for s in spans)
            x0, y0 = ln["bbox"][0], ln["bbox"][1]
            # PyMuPDF splits a visual line into several "lines" when spans are far apart
            # (e.g. "SDE Intern ........ May 2025 - Jul 2025"); keep them on one line.
            if lines and abs(lines[-1].y0 - y0) < 2.0:
                lines[-1].text = f"{lines[-1].text}  {text}".strip()
                lines[-1].bold = lines[-1].bold and bold
                continue
            lines.append(Line(text=text, page=pno, x0=x0, y0=y0, size=round(size, 1), bold=bool(bold)))
        if lines:
            x0, y0, x1, y1 = blk["bbox"]
            blocks.append(_Block(lines=lines, x0=x0, y0=y0, x1=x1, y1=y1))
    return blocks


def _drop_running_headers(pages: list[tuple[int, float, float, list[_Block]]]) -> None:
    if len(pages) < 2:
        return
    norm = lambda s: re.sub(r"\d+", "#", s.lower()).strip()  # noqa: E731
    counts: Counter[str] = Counter()
    for _pno, _w, h, blocks in pages:
        seen = {norm(ln.text) for b in blocks for ln in b.lines if ln.y0 < 0.07 * h or ln.y0 > 0.93 * h}
        counts.update(seen)
    repeated = {k for k, c in counts.items() if c >= 2}
    if not repeated:
        return
    for i, (_pno, _w, h, blocks) in enumerate(pages):
        for b in blocks:
            b.lines = [ln for ln in b.lines
                       if not (i > 0 and (ln.y0 < 0.07 * h or ln.y0 > 0.93 * h) and norm(ln.text) in repeated)
                       and not re.fullmatch(r"page\s*#(\s*(of|/)\s*#)?", norm(ln.text))]
        pages[i] = (pages[i][0], pages[i][1], h, [b for b in blocks if b.lines])


def detect_gutter(blocks: list[_Block], page_w: float) -> float | None:
    """x position of an empty vertical gutter between two text columns, if any."""
    narrow = [b for b in blocks if b.width < 0.62 * page_w]
    if len(narrow) < 4:
        return None
    nbins = int(page_w) // 4 + 1
    bins = [0.0] * nbins
    for b in narrow:
        for i in range(max(0, int(b.x0) // 4), min(nbins, int(b.x1) // 4 + 1)):
            bins[i] += b.height
    lo, hi = int(0.2 * nbins), int(0.8 * nbins)
    best: list[int] = []
    run: list[int] = []
    for i in range(lo, hi):
        if bins[i] == 0:
            run.append(i)
        else:
            if len(run) > len(best):
                best = run
            run = []
    if len(run) > len(best):
        best = run
    if not best or len(best) * 4 < 8:
        return None
    left = sum(b.height for b in narrow if b.x1 <= best[0] * 4 + 4)
    right = sum(b.height for b in narrow if b.x0 >= best[-1] * 4)
    if left == 0 or right == 0:
        return None
    return (best[0] + best[-1]) / 2 * 4.0


def _reading_order(blocks: list[_Block], width: float, height: float) -> tuple[list[_Block], bool]:
    blocks = sorted(blocks, key=lambda b: (round(b.y0), b.x0))
    gutter = detect_gutter(blocks, width)
    if gutter is None:
        return blocks, False
    spanning = [b for b in blocks if b.x0 < gutter < b.x1]
    cols = [b for b in blocks if not (b.x0 < gutter < b.x1)]
    if not cols:
        return blocks, False
    col_top = min(b.y0 for b in cols)
    col_bot = max(b.y1 for b in cols)
    above = [b for b in spanning if b.y1 <= col_top + 1]
    below = [b for b in spanning if b.y0 >= col_bot - 1]
    middle = [b for b in spanning if b not in above and b not in below]
    left = sorted([b for b in cols if b.x1 <= gutter] + middle, key=lambda b: b.y0)
    right = sorted([b for b in cols if b.x0 >= gutter], key=lambda b: b.y0)
    return above + left + right + below, True


# --- DOCX ---------------------------------------------------------------------

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _parse_docx(data: bytes) -> ParsedDocument:
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError("unreadable_document", f"Could not open the DOCX: {exc}") from exc

    lines: list[Line] = []
    seen_cells: set[int] = set()

    def para_line(p: Paragraph, region: str = "body") -> Line | None:
        text = clean_text(p.text)
        if not text:
            return None
        runs = [r for r in p.runs if r.text.strip()]
        bold = bool(runs) and all(bool(r.bold) or bool(p.style is not None and p.style.font.bold) for r in runs)
        sizes = [r.font.size.pt for r in runs if r.font.size is not None]
        style = p.style.name if p.style is not None else None
        # Word bullets are numbering properties, not characters
        if (p._p.pPr is not None and p._p.pPr.numPr is not None) or (style or "").lower().startswith("list"):
            text = text if text.startswith("- ") else "- " + text
        return Line(text=text, size=max(sizes) if sizes else 11.0, bold=bold, style=style, region=region)

    def walk_table(tbl: Table) -> None:
        for row in tbl.rows:
            for cell in row.cells:
                if id(cell._tc) in seen_cells:  # merged cells repeat the same element
                    continue
                seen_cells.add(id(cell._tc))
                for p in cell.paragraphs:
                    ln = para_line(p)
                    if ln:
                        lines.append(ln)
                for inner in cell.tables:
                    walk_table(inner)

    # Header first: contact details often live there.
    try:
        hdr = document.sections[0].header
        for p in hdr.paragraphs:
            ln = para_line(p, region="header")
            if ln:
                lines.append(ln)
    except (IndexError, AttributeError):
        pass

    seen_box_text: set[str] = set()
    for child in document.element.body.iterchildren():
        if child.tag == _W + "p":
            p = Paragraph(child, document)
            ln = para_line(p)
            if ln:
                lines.append(ln)
            # text boxes anchored in this paragraph (skip mc:Fallback duplicates)
            for tb in child.iter(_W + "txbxContent"):
                if any(a.tag.endswith("Fallback") for a in tb.iterancestors()):
                    continue
                for bp in tb.iter(_W + "p"):
                    t = clean_text("".join(x.text or "" for x in bp.iter(_W + "t")))
                    if t and t not in seen_box_text:
                        seen_box_text.add(t)
                        lines.append(Line(text=t))
        elif child.tag == _W + "tbl":
            walk_table(Table(child, document))
    return ParsedDocument(lines=lines, page_count=1, layout="docx")
